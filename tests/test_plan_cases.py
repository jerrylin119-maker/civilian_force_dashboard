# -*- coding: utf-8 -*-
"""活動類計畫類型、歷史案例範例與「油資、過路費、停車費不予補助」檢查的測試。"""

import io
import json
import re
import unittest
from datetime import date

import plan_cases
import plan_generator as pg


TODAY = date(2026, 10, 27)
CASE_KEYS = list(plan_cases.CASE_KEYS)


def _titles(findings, level=None):
    return [f.title for f in findings if level is None or f.level == level]


def _activity(**overrides):
    start = pg.earliest_activity_date(TODAY)
    subsidy = overrides.pop("subsidy", 20000)
    raw_items = overrides.pop("items", [
        {"name": "場地佈置", "kind": "一般", "qty": 1, "unit": "式", "price": 20000, "use": "供典禮使用"},
        {"name": "雜支", "kind": "雜支", "qty": 1, "unit": "式", "price": 1000, "use": "文具"},
        {"name": "便餐", "kind": "餐費", "qty": 2, "unit": "桌", "price": 5000, "use": "供用餐"},
    ])
    base = dict(
        plan_type="activity", unit="成功義消大隊", year=115, name="績優義消人員表揚",
        start=start, end=start, place="成功消防大隊車庫", audience="義消人員", headcount=100,
        method="", instructors="", co_organizer="", supervisor="", councilors="王議員",
        subsidy=subsidy, items=pg.normalize_items(raw_items, subsidy=subsidy), courses=[], roster=[],
    )
    base.update(overrides)
    return pg.PlanContext(**base)


class ActivityTypeTests(unittest.TestCase):
    def test_both_plan_types_registered_with_labels(self):
        self.assertIn("activity", pg.PLAN_TYPES)
        self.assertEqual(pg.PLAN_TYPES["training"]["audience_label"], "參訓對象")
        self.assertEqual(pg.PLAN_TYPES["activity"]["audience_label"], "參加對象")
        self.assertEqual(pg.MODULE_REVISION, 10)

    def test_activity_sections_use_activity_wording_and_skip_award(self):
        ctx = _activity(purpose="表彰績優人員", method="（一）績優表揚。\n（二）頒獎。")
        heads = [h for h, _ in pg.sections(ctx)]
        self.assertEqual(heads, ["目的", "辦理單位", "活動日期及地點", "參加對象", "活動內容", "活動經費", "其他"])
        self.assertNotIn("訓練課程", heads)
        self.assertNotIn("緣由", heads)
        text = pg.markdown_preview(ctx)
        self.assertNotIn("嘉獎", text)
        self.assertIn("臺東縣消防局辦理成功義消大隊", text)
        self.assertIn("115年度績優義消人員表揚活動計畫", text)

    def test_fund_text_points_to_first_attachment_when_no_schedule(self):
        ctx = _activity()
        self.assertIn("本案活動經費共計新臺幣", pg.fund_text(ctx))
        self.assertIn("（如附件一）", pg.fund_text(ctx))
        self.assertEqual(pg.attachment_labels(ctx), (None, "附件一", "附件二"))

    def test_training_attachments_unchanged(self):
        start = pg.earliest_activity_date(TODAY)
        ctx = pg.PlanContext(
            plan_type="training", unit="達仁義消分隊", year=115, name="繩索救助訓練", start=start, end=start,
            place="x", audience="y", headcount=1, method="", instructors="", co_organizer="", supervisor="",
            councilors="", subsidy=0, items=[], courses=[], roster=[])
        self.assertEqual(pg.attachment_labels(ctx), ("附件一", "附件二", "附件三"))
        self.assertIn("（如附件二）", pg.fund_text(ctx))

    def test_schedule_not_required_for_honor_but_required_for_study(self):
        honor = _activity()
        self.assertEqual(_titles(pg.check_plan(honor, today=TODAY), pg.ERROR), [])
        study = _activity(name="縣外研習")
        errors = _titles(pg.check_plan(study, today=TODAY), pg.ERROR)
        self.assertIn("沒有活動流程／行程表", errors)
        env = _activity(name="員工環境教育暨文康")
        self.assertIn("沒有活動流程／行程表", _titles(pg.check_plan(env, today=TODAY), pg.ERROR))

    def test_required_field_label_follows_plan_type(self):
        ctx = _activity(audience="")
        hit = [f for f in pg.check_plan(ctx, today=TODAY) if f.title == "尚有必填欄位未填"]
        self.assertIn("參加對象", hit[0].detail)
        self.assertNotIn("參訓對象", hit[0].detail)

    def test_docx_and_xlsx_without_schedule_have_no_course_sheet(self):
        from docx import Document
        from openpyxl import load_workbook
        ctx = _activity(method="（一）績優表揚。")
        doc = Document(io.BytesIO(pg.build_docx(ctx)))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("附件一", text)
        self.assertNotIn("附件二", text)
        self.assertNotIn("課程表", text)
        wb = load_workbook(io.BytesIO(pg.build_xlsx(ctx)))
        self.assertEqual(wb.sheetnames, ["經費概算表"])

    def test_docx_and_xlsx_with_schedule_use_activity_names(self):
        from docx import Document
        from openpyxl import load_workbook
        start = pg.earliest_activity_date(TODAY)
        ctx = _activity(name="縣外研習", courses=[{"date": start, "time": "09:00", "content": "研習交流", "teacher": ""}],
                        roster=[{"unit": "成功義消大隊", "title": "隊員", "name": "甲"}])
        doc = Document(io.BytesIO(pg.build_docx(ctx)))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("縣外研習活動流程表", text)
        self.assertIn("附件三", text)
        self.assertIn("參加人員名冊", text)
        self.assertNotIn("訓練課程表", text)
        wb = load_workbook(io.BytesIO(pg.build_xlsx(ctx)))
        self.assertEqual(wb.sheetnames, ["經費概算表", "活動流程表"])


class NoSubsidyWordsTests(unittest.TestCase):
    def test_fuel_toll_parking_items_get_no_subsidy_and_warn(self):
        items = pg.normalize_items([
            {"name": "遊覽車租金", "kind": "一般", "qty": 2, "unit": "日", "price": 15000, "use": "往返交通"},
            {"name": "油資", "kind": "一般", "qty": 1, "unit": "式", "price": 5000, "use": "遊覽車用油"},
            {"name": "過路費", "kind": "一般", "qty": 1, "unit": "式", "price": 1000, "use": "國道通行"},
            {"name": "停車費", "kind": "一般", "qty": 1, "unit": "式", "price": 500, "use": "景點停車"},
        ], subsidy=20000)
        by = {i["name"]: i for i in items}
        self.assertEqual(by["遊覽車租金"]["subsidy"], 20000)
        for name in ("油資", "過路費", "停車費"):
            self.assertEqual(by[name]["subsidy"], 0)
            self.assertEqual(by[name]["self"], by[name]["total"])
            self.assertNotIn("議員建議補助款", by[name]["remark"])
        ctx = _activity(items=[{"name": i["name"], "kind": i["kind"], "qty": i["qty"], "unit": i["unit"],
                                "price": i["price"], "use": i["use"]} for i in items], subsidy=20000)
        warn = [f for f in pg.check_plan(ctx, today=TODAY) if "油資、過路費或停車費" in f.title]
        self.assertEqual(len(warn), 1)
        self.assertEqual(warn[0].level, pg.WARN)
        self.assertIn("油資", warn[0].detail)
        self.assertIn("過路費", warn[0].detail)
        self.assertIn("停車費", warn[0].detail)

    def test_bundled_name_is_flagged_with_split_advice(self):
        ctx = _activity(items=[
            {"name": "遊覽車租金（2日，含油資、過路費、駕駛食宿）", "kind": "一般", "qty": 2, "unit": "日",
             "price": 15000, "use": "往返交通"},
            {"name": "住宿", "kind": "一般", "qty": 20, "unit": "間", "price": 3000, "use": "住宿"},
        ], subsidy=60000)
        warn = [f for f in pg.check_plan(ctx, today=TODAY) if "油資、過路費或停車費" in f.title]
        self.assertEqual(len(warn), 1)
        self.assertIn("拆成", warn[0].detail)
        self.assertEqual(ctx.items[0]["subsidy"], 0)
        self.assertEqual(ctx.items[1]["subsidy"], 60000)

    def test_no_warning_without_those_words(self):
        ctx = _activity()
        self.assertEqual([f for f in pg.check_plan(ctx, today=TODAY) if "油資" in f.title], [])


class HistoricalCaseTests(unittest.TestCase):
    def test_cases_listed_by_plan_type(self):
        activity = [k for k, _l, _s in pg.case_list("activity")]
        training = [k for k, _l, _s in pg.case_list("training")]
        self.assertEqual(activity, ["honor", "study_tour", "culture_edu", "award_dinner", "env_study"])
        self.assertEqual(training, [
            "vehicle_fire", "rehab_bls", "mountain_rescue", "car_rescue",
            "rope_two_days", "water_rescue", "mountain_three_days",
        ])
        equipment = [k for k, _l, _s in pg.case_list("equipment")]
        self.assertEqual(equipment, ["equip_rescue"])
        self.assertEqual(activity + training + equipment, CASE_KEYS)

    def test_original_totals_are_preserved(self):
        expected = {
            "honor": (346000, 240000), "study_tour": (144000, 100000), "culture_edu": (78200, 40000),
            "vehicle_fire": (24500, 20000), "rehab_bls": (78600, 50000), "mountain_rescue": (20160, 20000),
            "car_rescue": (21300, 20000), "rope_two_days": (23900, 20000), "water_rescue": (30500, 20000),
            "mountain_three_days": (86750, 60000),
            "award_dinner": (48000, 20000), "env_study": (20000, 20000),
            "equip_rescue": (97423, 97423),
        }
        for key, (total, subsidy) in expected.items():
            ctx = pg.case_context(key, TODAY)
            self.assertEqual(ctx.total, total, key)
            self.assertEqual(ctx.subsidy, subsidy, key)
            self.assertEqual(sum(i["subsidy"] for i in ctx.items), subsidy, key)

    def test_each_case_passes_without_errors_and_exports(self):
        for key in CASE_KEYS:
            ctx = pg.case_context(key, TODAY)
            errors = _titles(pg.check_plan(ctx, today=TODAY), pg.ERROR)
            self.assertEqual(errors, [], key)
            self.assertGreater(len(pg.build_docx(ctx)), 1000)
            self.assertGreater(len(pg.build_xlsx(ctx)), 1000)
            self.assertEqual(ctx.plan_type, pg.case_inputs(key, TODAY)["plan_type"])

    def test_training_cases_have_schedule_and_no_award(self):
        for key in ("vehicle_fire", "rehab_bls", "mountain_rescue", "car_rescue",
                    "rope_two_days", "water_rescue", "mountain_three_days"):
            ctx = pg.case_context(key, TODAY)
            self.assertEqual(ctx.plan_type, "training", key)
            self.assertGreaterEqual(len(ctx.courses), 4, key)
            self.assertNotIn("沒有課程表", _titles(pg.check_plan(ctx, today=TODAY)), key)


    def test_honor_case_over_150k_gets_procurement_warning_only(self):
        ctx = pg.case_context("honor", TODAY)
        findings = pg.check_plan(ctx, today=TODAY)
        self.assertIn("總經費超過 15 萬元", _titles(findings, pg.WARN))
        self.assertEqual(ctx.courses, [])

    def test_study_and_culture_cases_have_schedules(self):
        self.assertGreaterEqual(len(pg.case_context("study_tour", TODAY).courses), 6)
        self.assertGreaterEqual(len(pg.case_context("culture_edu", TODAY).courses), 6)

    def test_study_case_has_no_bundled_fuel_toll(self):
        ctx = pg.case_context("study_tour", TODAY)
        self.assertEqual([i for i in ctx.items if pg._no_subsidy_words(i)], [])
        self.assertEqual([f for f in pg.check_plan(ctx, today=TODAY) if "油資" in f.title], [])

    def test_dates_are_relative_to_today(self):
        data = pg.case_inputs("study_tour", TODAY)
        self.assertGreaterEqual(data["start"], pg.earliest_activity_date(TODAY))
        self.assertEqual((data["end"] - data["start"]).days, 1)
        self.assertEqual(data["year"], TODAY.year - 1911)

    def test_cases_contain_no_personal_data(self):
        blob = ""
        for key in CASE_KEYS:
            data = pg.case_inputs(key, TODAY)
            data = {k: (v.isoformat() if isinstance(v, date) else v) for k, v in data.items()
                    if k not in ("items", "courses")}
            blob += json.dumps(data, ensure_ascii=False)
            blob += json.dumps([[i["name"], i["use"]] for i in pg.case_inputs(key, TODAY)["items"]], ensure_ascii=False)
            blob += json.dumps([[c["content"], c["teacher"]] for c in pg.case_inputs(key, TODAY)["courses"]],
                               ensure_ascii=False)
            self.assertEqual(pg.case_context(key, TODAY).roster, [])
        self.assertIsNone(re.search(r"[A-Z][12]\d{8}", blob), "不得含身分證字號")
        self.assertIsNone(re.search(r"0\d{1,2}-?\d{6,8}", blob), "不得含電話")
        self.assertIsNone(re.search(r"\d{6,}", blob.replace("000", "")), "不得含案號、發文字號")
        for word in ("廖昱翔", "林威宇", "林顯峯", "盧東發", "管建興", "吳秀華", "高美珠", "楊招信"):
            self.assertNotIn(word, blob)
        self.assertNotIn("嘉獎", blob)

    def test_fresh_dict_each_call(self):
        a = pg.case_inputs("honor", TODAY)
        a["items"].append({"name": "x"})
        self.assertEqual(len(pg.case_inputs("honor", TODAY)["items"]), 8)

    def test_submission_form_same_wording_for_activity_and_training(self):
        """陳報單句型相同，只有計畫名稱與檢陳文件不同；預訂日期地點寫在第一點；不得出現敘獎額度。"""
        for key in CASE_KEYS:
            ctx = pg.case_context(key, TODAY)
            subject = pg.submission_subject(ctx)
            body = pg.submission_body(ctx)
            joined = "".join(body)
            self.assertIn(pg.plan_full_name(ctx), subject)
            self.assertIn("乙案，請鑒核", subject)
            self.assertIn("「" + pg.plan_full_name(ctx) + "」", subject)
            if ctx.plan_type != "equipment":
                self.assertIn(ctx.date_text, body[0])
                self.assertIn(ctx.place.strip(), body[0])
            self.assertIn("建請鈞長同意辦理申請補助款事宜", body[1])
            self.assertIn(f"新臺幣{pg.to_cn_amount(ctx.total)}元整", body[1])
            self.assertIn("檢陳相關資料附件如下", body[2])
            self.assertIn("計畫", body[2])
            self.assertNotIn("嘉獎", joined)
            self.assertNotIn("人事室", pg.meet_text(ctx))
            self.assertGreater(len(pg.build_submission_docx(ctx)), 2000)

    def test_submission_form_without_self_fund_omits_rest_clause(self):
        ctx = pg.case_context("env_study", TODAY)
        self.assertLess(ctx.self_fund, 0.5)
        self.assertNotIn("其餘", "".join(pg.submission_body(ctx)))
        self.assertNotIn("不足部分", pg.fund_text(ctx))
        ctx = pg.case_context("honor", TODAY)
        self.assertIn("其餘新臺幣", "".join(pg.submission_body(ctx)))
        self.assertIn("自籌款支應", pg.submission_body(ctx)[1])

    def test_training_submission_lists_safety_checklist_and_default_meet(self):
        ctx = pg.case_context("rehab_bls", TODAY)
        self.assertIn("訓練安全計畫與檢核表", pg.submission_body(ctx)[2])
        self.assertEqual(pg.meet_text(ctx), "災害搶救科、綜合企劃及安全衛生科、民力運用及訓練科、會計室")


class MultiSourceAndEquipmentTests(unittest.TestCase):
    def _mixed(self):
        data = pg.case_inputs("rehab_bls", TODAY)
        data["sources"] = [
            {"kind": "councilor", "names": "林", "amount": 30000},
            {"kind": "rep", "names": "陳", "township": "卑南鄉公所", "amount": 20000},
        ]
        data["awards"] = [{"role": "大隊督辦人員", "count": 2, "kind": "嘉獎", "times": 1},
                          {"role": "分隊承辦人員", "count": 1, "kind": "嘉獎", "times": 2}]
        return pg.context_from_inputs(data, "training")

    def test_sources_sum_to_subsidy_and_each_item_remark_names_every_source(self):
        ctx = self._mixed()
        self.assertEqual(ctx.subsidy, 50000)
        self.assertEqual(sum(i["subsidy"] for i in ctx.items), 50000)
        for item in ctx.items:
            self.assertAlmostEqual(sum(a for _l, a in item["by_source"]), item["subsidy"], places=6)
            self.assertNotIn("支應", item["remark"])             # 備註不顯示各來源金額
            self.assertNotIn("元由", item["remark"])
        by_label = {}
        for item in ctx.items:
            for label, amount in item["by_source"]:
                by_label[label] = by_label.get(label, 0) + amount
        self.assertEqual(by_label, {"議員建議補助款": 30000, "卑南鄉公所代表建議補助款": 20000})

    def test_mixed_source_texts_and_attachments(self):
        ctx = self._mixed()
        fund = pg.fund_text(ctx)
        self.assertIn("林議員建議縣府補助新臺幣參萬元整", fund)
        self.assertIn("卑南鄉公所陳代表建議卑南鄉公所補助新臺幣貳萬元整", fund)
        body = pg.submission_body(ctx)
        self.assertIn("新臺幣參萬元整由本縣縣議會林議員建議本縣縣府補助經費項下支應", body[1])
        self.assertIn("新臺幣貳萬元整由卑南鄉公所陳代表建議卑南鄉公所補助經費項下支應", body[1])
        self.assertIn("臺東縣議員建議補助活動類案件建議表1份", body[2])
        self.assertIn("卑南鄉代表建議補助活動類案件建議表1份", body[2])
        self.assertTrue(pg.submission_subject(ctx).endswith("乙案，請鑒核。"))

    def test_single_source_keeps_old_wording(self):
        ctx = pg.case_context("honor", TODAY)
        self.assertEqual(len(ctx.sources), 1)
        self.assertIn("王議員、李議員建議縣府補助新臺幣貳拾肆萬元整", pg.fund_text(ctx))
        for item in ctx.items:
            self.assertNotIn("代表", item["remark"])

    def test_rep_only_source_goes_through_battalion(self):
        data = pg.case_inputs("honor", TODAY)
        data["sources"] = [{"kind": "rep", "names": "", "township": "", "amount": 240000}]
        ctx = pg.context_from_inputs(data, "activity")
        self.assertIn("ＯＯ公所", pg.fund_text(ctx))
        self.assertIn("由大隊發文", pg.issuing_note(ctx))
        self.assertIn("縣府", pg.issuing_note(pg.case_context("honor", TODAY)))

    def test_awards_add_clause_and_personnel_office(self):
        ctx = self._mixed()
        body = pg.submission_body(ctx)
        self.assertTrue(body[3].startswith("四、獎懲：辦理本案出（不）力人員依規定獎（懲）勵，敘獎額度如下："))
        self.assertIn("（一）大隊督辦人員（2名）嘉獎1次。", body[3])
        self.assertIn("（二）分隊承辦人員（1名）嘉獎2次。", body[3])
        self.assertIn("人事室", pg.meet_text(ctx))
        self.assertNotIn("嘉獎", pg.markdown_preview(ctx))
        self.assertEqual(pg.normalize_awards([{"role": "", "count": 1, "times": 1}, {"role": "x", "count": 0, "times": 1}]), [])

    def test_equipment_type_sections_and_rules(self):
        ctx = pg.case_context("equip_rescue", TODAY)
        self.assertEqual(ctx.title2, "115年度購置救災設備及後勤照護設備計畫")
        heads = [h for h, _ in pg.sections(ctx)]
        self.assertEqual(heads, ["緣由", "目的", "辦理單位", "設置地點", "使用單位", "購置項目", "預計完成日期",
                                 "購置經費", "驗收及管理", "其他"])
        findings = pg.check_plan(ctx, today=TODAY)
        titles = [f.title for f in findings]
        self.assertIn("設備類不適用「自籌款 30%」規定", titles)
        self.assertEqual(_titles(findings, pg.ERROR), [])
        self.assertNotIn("餐費及紀念品未超過補助金額 50%", titles)
        self.assertNotIn("沒有課程表", titles)
        body = pg.submission_body(ctx)
        self.assertIn("估價單2份", body[2])
        self.assertIn("審核表1份", body[2])
        self.assertIn("臺東縣議員建議補助設備類案件建議表1份", body[2])
        self.assertEqual(pg.meet_text(ctx), "災害搶救科")

    def test_equipment_allows_equipment_kind_but_activity_still_rejects_it(self):
        ctx = pg.case_context("equip_rescue", TODAY)
        self.assertNotIn("活動類補助款不得用於資本門（購置設備）", _titles(pg.check_plan(ctx, today=TODAY), pg.ERROR))
        data = pg.case_inputs("honor", TODAY)
        data["items"].append({"name": "發電機", "kind": "設備購置", "qty": 1, "unit": "台", "price": 1000, "use": "x"})
        act = pg.context_from_inputs(data, "activity")
        self.assertIn("活動類補助款不得用於資本門（購置設備）", _titles(pg.check_plan(act, today=TODAY), pg.ERROR))

    def test_equipment_warns_on_personal_gear_and_meals(self):
        data = pg.case_inputs("equip_rescue", TODAY)
        data["items"] += [
            {"name": "個人制服", "kind": "設備購置", "qty": 1, "unit": "件", "price": 1000, "use": "x"},
            {"name": "便當", "kind": "餐費", "qty": 1, "unit": "份", "price": 100, "use": "x"},
        ]
        titles = _titles(pg.check_plan(pg.context_from_inputs(data, "equipment"), today=TODAY), pg.WARN)
        self.assertIn("設備類不得採購個人服裝及裝備", titles)
        self.assertIn("設備類概算表出現餐費或紀念品", titles)

    def test_equipment_docx_and_xlsx_have_no_schedule(self):
        from docx import Document
        from openpyxl import load_workbook
        ctx = pg.case_context("equip_rescue", TODAY)
        text = "\n".join(p.text for p in Document(io.BytesIO(pg.build_docx(ctx))).paragraphs)
        self.assertIn("購置項目", text)
        self.assertNotIn("課程表", text)
        self.assertEqual(load_workbook(io.BytesIO(pg.build_xlsx(ctx))).sheetnames, ["經費概算表"])


class SubmissionFormatTests(unittest.TestCase):
    """陳報單格式（比照臺東分隊 115 年設備案實際陳報單）與獎懲標準。"""

    def _equip(self):
        ctx = pg.case_context("equip_rescue", TODAY)
        ctx.org_battalion, ctx.org_unit, ctx.unit = "台東大隊", "台東分隊", "臺東義消分隊"
        ctx.sources = pg.normalize_sources([{"kind": "councilor", "names": "陳志峰", "amount": ctx.subsidy}])
        ctx.award_basis = "amount"
        return ctx

    def test_matches_real_sample_sentences(self):
        ctx = self._equip()
        self.assertEqual(pg.org_title(ctx), "臺東縣消防局臺東大隊臺東分隊")
        self.assertEqual(pg.submission_subject(ctx),
                         "有關本分隊辦理「臺東義消分隊115年度購置救災設備及後勤照護設備計畫」乙案，請鑒核。")
        body = pg.submission_body(ctx)
        self.assertEqual(body[1], "二、本案所需經費新臺幣玖萬柒仟肆佰貳拾參元整，由本縣縣議會陳議員志峰"
                                  "建議本縣縣府補助經費項下支應，建請鈞長同意辦理申請補助款事宜。")
        self.assertEqual(body[2].split("\n"), [
            "三、檢陳相關資料附件如下：",
            "　（一）115年度購置救災設備及後勤照護設備計畫1份。",
            "　（二）臺東縣議員建議補助設備類案件建議表1份。",
            "　（三）審核表1份。",
            "　（四）估價單2份。"])
        self.assertEqual(body[3], "四、獎懲：依本縣消防局消防專業人員平時考核獎懲標準規定，爭取建議補助款總額度達五萬元至十萬元以下者"
                                  "單位主管及承辦人嘉獎1次；協辦人員(至多2名)嘉獎1次。")
        self.assertEqual(pg.meet_text(ctx), "災害搶救科、人事室")
        self.assertGreater(len(pg.build_submission_docx(ctx)), 2000)

    def test_battalion_account_header_and_wording(self):
        ctx = self._equip()
        ctx.org_unit = ""
        self.assertEqual(pg.org_title(ctx), "臺東縣消防局臺東大隊")
        self.assertIn("有關本大隊辦理", pg.submission_subject(ctx))
        ctx.org_battalion = ""
        self.assertEqual(pg.org_title(ctx), "臺東縣消防局○○大隊○○分隊")

    def test_councilor_call(self):
        self.assertEqual(pg._councilor_call("陳志峰"), "陳議員志峰")
        self.assertEqual(pg._councilor_call("陳志峰議員、王明"), "陳議員志峰、王議員明")
        self.assertEqual(pg._councilor_call("歐陽志峰"), "歐陽議員志峰")
        self.assertEqual(pg._councilor_call(""), "ＯＯ議員")
        self.assertEqual(pg._councilor_call("Kacaw"), "Kacaw議員")

    def test_award_by_amount_tiers(self):
        cases = {19999: "", 20000: "達兩萬元以上，未達五萬元者承辦人嘉獎1次；協辦人員(至多2名)嘉獎1次", 49999: "達兩萬元以上，未達五萬元",
                 50000: "達五萬元至十萬元以下者", 99999: "達五萬元至十萬元以下者",
                 100000: "達十萬元至三十萬元以下者單位主管及承辦人嘉獎2次；協辦人員(至多2名)嘉獎2次",
                 299999: "達十萬元至三十萬元以下者", 300000: "達三十萬元以上者單位主管及承辦人記功1次；協辦人員(至多2名)嘉獎2次"}
        for amount, expect in cases.items():
            got = pg.award_by_amount(amount)
            if expect:
                self.assertIn(expect, got, amount)
                self.assertTrue(got.startswith("依本縣消防局消防專業人員平時考核獎懲標準規定，"))
            else:
                self.assertEqual(got, "")

    def test_award_by_days(self):
        d = date(2026, 12, 1)
        self.assertIn("達三日以上者承辦人嘉獎2次；單位主管嘉獎2次；協辦人員嘉獎1次", pg.award_by_days(d, date(2026, 12, 3)))
        self.assertIn("在二日內者承辦人嘉獎2次；單位主管嘉獎1次；協辦人員嘉獎1次", pg.award_by_days(d, date(2026, 12, 2)))
        self.assertIn("在一日內者承辦人嘉獎2次；單位主管嘉獎1次；協辦人員嘉獎1次", pg.award_by_days(d, d))
        self.assertEqual(pg.award_by_days(None, d), "")
        self.assertEqual(pg.award_by_days(d, date(2026, 11, 1)), "")

    def test_award_basis_switches_and_adds_personnel_office(self):
        ctx = pg.case_context("honor", TODAY)                      # 補助 24 萬，單日活動
        self.assertEqual(pg.award_text(ctx), "")                    # 預設不敘獎
        self.assertNotIn("人事室", pg.meet_text(ctx))
        ctx.award_basis = "amount"
        self.assertIn("達十萬元至三十萬元以下者", pg.award_text(ctx))
        self.assertIn("人事室", pg.meet_text(ctx))
        self.assertTrue(pg.submission_body(ctx)[-1].startswith("四、獎懲："))
        ctx.award_basis = "days"
        self.assertIn("在一日內者", pg.award_text(ctx))
        ctx.award_basis = "none"
        self.assertEqual(pg.award_text(ctx), "")
        self.assertEqual(len(pg.submission_body(ctx)), 3)

    def test_equipment_has_no_day_based_award(self):
        ctx = self._equip()
        ctx.award_basis = "days"
        self.assertEqual(pg.award_text(ctx), "")

    def test_self_fund_wording_uses_chinese_capitals(self):
        ctx = pg.case_context("honor", TODAY)
        text = pg.submission_body(ctx)[1]
        self.assertIn(f"新臺幣{pg.to_cn_amount(ctx.total)}元整", text)
        self.assertIn(f"其餘新臺幣{pg.to_cn_amount(ctx.self_fund)}元整", text)


class QuantityAndCaseNumberTests(unittest.TestCase):
    """名冊（參加）人數與便當、飲品、耗材數量要合理對應；案號會印在文件上。"""

    def _ctx(self, items, headcount=30, roster=None, plan_type="training", case="rehab_bls"):
        data = pg.case_inputs(case, TODAY)
        data["headcount"] = headcount
        data["items"] = items
        ctx = pg.context_from_inputs(data, plan_type)
        if roster is not None:
            ctx.roster = roster
        return ctx

    def _meal(self, qty, use, name="便當"):
        return [{"name": name, "kind": "餐費", "qty": qty, "unit": "份", "price": 100, "use": use},
                {"name": "紅布條", "kind": "一般", "qty": 1, "unit": "條", "price": 1000, "use": "背景布條"}]

    def _titles(self, ctx, level):
        return [f.title for f in pg.check_plan(ctx, today=TODAY) if f.level == level]

    def test_meal_far_above_headcount_is_an_error_without_explanation(self):
        ctx = self._ctx(self._meal(100, "供訓練人員用餐"), headcount=30)
        self.assertIn("「便當」數量與人數不相符", self._titles(ctx, pg.ERROR))

    def test_meal_with_quantity_source_in_note_passes(self):
        ctx = self._ctx(self._meal(100, "學員30人＋工作人員10人＋來賓60人，午餐及晚餐"), headcount=30)
        self.assertNotIn("「便當」數量與人數不相符", self._titles(ctx, pg.ERROR))
        self.assertIn("「便當」數量超過人數基準，已於用途說明註明來源", self._titles(ctx, pg.OK))

    def test_lunch_plus_dinner_is_two_meals_per_person(self):
        for name, use in (("午晚餐便當", "用餐"), ("便當", "供午餐及晚餐"), ("便當", "午、晚餐"), ("便當", "供訓練人員中餐及晚餐")):
            ctx = self._ctx(self._meal(100, use, name), headcount=50)
            self.assertNotIn("「%s」數量與人數不相符" % name, self._titles(ctx, pg.ERROR), (name, use))
        ctx = self._ctx(self._meal(100, "午餐"), headcount=50)               # 只有午餐，100 份就偏多
        self.assertIn("「便當」數量與人數不相符", self._titles(ctx, pg.ERROR))
        self.assertEqual(pg._meals_in("午晚餐"), (2, "午餐＋晚餐"))
        self.assertEqual(pg._meals_in("早餐、午餐及晚餐")[0], 3)
        self.assertEqual(pg._meals_in("便當")[0], 1)

    def test_roster_count_takes_priority_over_headcount(self):
        roster = [{"unit": "x", "title": "y", "name": f"人{n}"} for n in range(10)]
        ctx = self._ctx(self._meal(30, "用餐"), headcount=50, roster=roster)
        self.assertIn("「便當」數量與人數不相符", self._titles(ctx, pg.ERROR))
        self.assertIn("名冊 10 人", [f.detail for f in pg.check_plan(ctx, today=TODAY)
                                   if f.title == "「便當」數量與人數不相符"][0])

    def test_multi_day_scales_the_allowance(self):
        data = pg.case_inputs("rehab_bls", TODAY)
        data["start"], data["end"] = date(2026, 12, 11), date(2026, 12, 12)
        data["headcount"] = 30
        data["items"] = self._meal(60, "午餐")
        ctx = pg.context_from_inputs(data, "training")
        self.assertEqual(ctx.days, 2)
        self.assertNotIn("「便當」數量與人數不相符", self._titles(ctx, pg.ERROR))

    def test_drinks_and_goods_are_warnings_not_errors(self):
        items = [{"name": "礦泉水", "kind": "一般", "qty": 300, "unit": "瓶", "price": 10, "use": "飲用"},
                 {"name": "毛巾", "kind": "一般", "qty": 100, "unit": "條", "price": 30, "use": "發給學員"}]
        ctx = self._ctx(items, headcount=30)
        warns = self._titles(ctx, pg.WARN)
        self.assertIn("「礦泉水」數量與人數不相符", warns)
        self.assertIn("「毛巾」數量與人數不相符", warns)
        self.assertNotIn("「礦泉水」數量與人數不相符", self._titles(ctx, pg.ERROR))

    def test_reasonable_quantities_get_an_ok_line(self):
        ctx = self._ctx(self._meal(30, "供訓練人員午餐"), headcount=30)
        self.assertIn("餐點與耗材數量與人數相符", self._titles(ctx, pg.OK))

    def test_any_meal_above_headcount_gets_a_reminder(self):
        use = "供學員午餐"
        for qty, expect in ((30, None), (31, "「便當」數量比人數基準多"), (33, "「便當」數量比人數基準多"),
                            (34, "「便當」數量與人數不相符"), (100, "「便當」數量與人數不相符")):
            ctx = self._ctx(self._meal(qty, use), headcount=30)
            found = [f for f in pg.check_plan(ctx, today=TODAY) if f.title.startswith("「便當」")]
            if expect is None:
                self.assertEqual(found, [], qty)
            else:
                self.assertEqual([f.title for f in found], [expect], qty)
        mild = self._ctx(self._meal(33, use), headcount=30)
        self.assertFalse([t for t in self._titles(mild, pg.ERROR) if t.startswith("「便當」")])   # 只多一點：提醒，不擋下
        self.assertIn("「便當」數量比人數基準多", self._titles(mild, pg.WARN))

    def test_note_with_people_count_clears_a_mild_overage_too(self):
        ctx = self._ctx(self._meal(33, "學員30人＋工作人員3人，午餐"), headcount=30)
        titles = [f.title for f in pg.check_plan(ctx, today=TODAY)]
        self.assertNotIn("「便當」數量比人數基準多", titles)
        self.assertIn("「便當」數量超過人數基準，已於用途說明註明來源", titles)

    def test_table_meals_are_checked_with_a_looser_rule(self):
        def table(qty, headcount):
            items = [{"name": "便餐", "kind": "餐費", "qty": qty, "unit": "桌", "price": 4000, "use": "晚餐"},
                     {"name": "紅布條", "kind": "一般", "qty": 1, "unit": "條", "price": 1000, "use": "布條"}]
            ctx = self._ctx(items, headcount=headcount)
            return [(f.level, f.title) for f in pg.check_plan(ctx, today=TODAY) if f.title.startswith("「便餐」")]
        self.assertEqual(table(3, 25), [])                                    # 25 人 3 桌，正常
        self.assertEqual(table(7, 30), [(pg.WARN, "「便餐」數量比人數基準多")])   # 30 人 7 桌：提醒
        self.assertEqual(table(30, 25), [(pg.ERROR, "「便餐」數量與人數不相符")])  # 25 人 30 桌：離譜

    def test_meal_use_options_follow_the_plan_dates(self):
        d1, d2 = date(2026, 12, 12), date(2026, 12, 13)
        self.assertEqual(pg.meal_use_options(d1, d1),
                         ["115年12月12日中餐", "115年12月12日晚餐", "115年12月12日中餐及晚餐"])
        opts = pg.meal_use_options(d1, d2)
        self.assertEqual(len(opts), 6)
        self.assertIn("115年12月13日晚餐", opts)
        self.assertEqual(pg.meal_use_options(None, None), ["中餐", "晚餐", "中餐及晚餐"])
        self.assertEqual(pg.meal_use_options(d1, date(2027, 6, 1)), ["中餐", "晚餐", "中餐及晚餐"])   # 期間太長不展開

    def test_meal_rows_are_forced_into_the_dropdown_choices(self):
        d = date(2026, 12, 12)
        rows = [
            {"項目": "午餐", "類別": "餐費", "數量": 25, "單位": "份", "單價": 100, "用途說明": "25人午餐"},
            {"項目": "晚餐便餐", "類別": "餐費", "數量": 3, "單位": "桌", "單價": 3500, "用途說明": "供午餐及晚餐"},
            {"項目": "便當", "類別": "餐費", "數量": 30, "單位": "", "單價": 100, "用途說明": "115年12月12日晚餐"},
            {"項目": "紅布條", "類別": "一般", "數量": 1, "單位": "條", "單價": 1000, "用途說明": "背景"},
            {"項目": "", "類別": "餐費", "數量": 0, "單位": "", "單價": 0, "用途說明": ""},
        ]
        out = pg.normalize_meal_rows(rows, d, d)
        self.assertEqual((out[0]["項目"], out[0]["用途說明"], out[0]["數量說明"]), ("便當", "115年12月12日中餐", "25人午餐"))
        self.assertEqual((out[1]["項目"], out[1]["用途說明"]), ("便餐", "115年12月12日中餐及晚餐"))
        self.assertEqual((out[2]["用途說明"], out[2]["單位"]), ("115年12月12日晚餐", "份"))   # 已是選項：不動；空單位補預設
        self.assertEqual((out[3]["項目"], out[3]["用途說明"]), ("紅布條", "背景"))               # 非餐費原樣
        self.assertEqual(out[4]["用途說明"], "")                                                  # 空白列不亂填
        for r in out[:3]:
            self.assertIn(r["項目"], pg.MEAL_NAMES)
            self.assertIn(r["用途說明"], pg.meal_use_options(d, d))

    def test_meal_use_and_note_are_merged_into_the_remark(self):
        self.assertEqual(pg.merge_meal_use("115年12月12日中餐", ""), "115年12月12日中餐")
        self.assertEqual(pg.merge_meal_use("115年12月12日中餐", "學員30人＋工作人員5人"),
                         "115年12月12日中餐；學員30人＋工作人員5人")
        self.assertEqual(pg.merge_meal_use("", "說明"), "說明")

    def test_dated_meal_row_counts_only_that_day(self):
        data = pg.case_inputs("rehab_bls", TODAY)
        data["start"], data["end"] = date(2026, 12, 11), date(2026, 12, 12)
        data["headcount"] = 30
        data["items"] = self._meal(40, "115年12月11日中餐")          # 2 天活動，但這列只算 12/11 一天：30 人 → 超過
        ctx = pg.context_from_inputs(data, "training")
        titles = [f.title for f in pg.check_plan(ctx, today=TODAY)]
        self.assertTrue(any(t.startswith("「便當」數量") for t in titles))
        data["items"] = self._meal(40, "午餐")                      # 沒指定日期：整個期間 2 天，40 份合理
        ctx = pg.context_from_inputs(data, "training")
        self.assertFalse([f for f in pg.check_plan(ctx, today=TODAY) if f.title.startswith("「便當」")])

    def test_equipment_is_not_checked(self):
        ctx = pg.case_context("equip_rescue", TODAY)
        ctx.headcount = 1
        self.assertEqual(pg._quantity_findings(ctx), [])

    def test_remark_has_no_source_or_amount_text(self):
        ctx = pg.case_context("honor", TODAY)
        for item in ctx.items:
            self.assertNotIn("元由", item["remark"])
            self.assertNotIn("支應", item["remark"])
            self.assertEqual(item["remark"], item["use"].strip())

    def test_case_number_printed_on_all_documents(self):
        from docx import Document
        from openpyxl import load_workbook
        ctx = pg.case_context("honor", TODAY)
        for build in (pg.build_docx, pg.build_submission_docx):
            self.assertEqual(Document(io.BytesIO(build(ctx))).sections[0].header.paragraphs[0].text, "")
        ctx.case_no = "115-成-0001"
        for build in (pg.build_docx, pg.build_submission_docx):
            doc = Document(io.BytesIO(build(ctx)))
            self.assertEqual(doc.sections[0].header.paragraphs[0].text, "案件編號：115-成-0001")
        wb = load_workbook(io.BytesIO(pg.build_xlsx(ctx)))
        self.assertEqual(wb.active.oddHeader.right.text, "案件編號：115-成-0001")
        self.assertIn("案件編號：115-成-0001", pg.submission_markdown(ctx))


if __name__ == "__main__":
    unittest.main()
