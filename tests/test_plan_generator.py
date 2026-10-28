# -*- coding: utf-8 -*-
"""計畫書／概算表產生器的規則與產出測試。"""

import io
import unittest
from datetime import date, timedelta

import plan_generator as pg


TODAY = date(2026, 10, 22)


def _items(rows, subsidy):
    return pg.normalize_items(rows, subsidy=subsidy)


def _ctx(today=TODAY, **overrides):
    start = overrides.pop("start", pg.earliest_activity_date(today))
    end = overrides.pop("end", start)
    subsidy = overrides.pop("subsidy", 20000)
    raw_items = overrides.pop("items", [
        {"name": "便當", "kind": "餐費", "qty": 60, "unit": "份", "price": 100, "use": "供訓練人員中餐及晚餐"},
        {"name": "棚架", "kind": "一般", "qty": 2, "unit": "式", "price": 3500, "use": "供訓練人員休息使用"},
        {"name": "礦泉水", "kind": "一般", "qty": 10, "unit": "箱", "price": 150, "use": "供訓練人員飲用"},
        {"name": "救助吊帶", "kind": "一般", "qty": 5, "unit": "條", "price": 2200, "use": "供繩索救援確保訓練人員安全"},
    ])
    courses = overrides.pop("courses", [
        {"date": start, "time": "09:10－12:00", "content": "繩索救助基礎操作", "teacher": "訓練教官"},
    ])
    base = dict(
        plan_type="training", unit="達仁義消分隊", year=115, name="繩索救助訓練",
        start=start, end=end, place="達仁分隊訓練廣場", audience="義消人員", headcount=30,
        method="", instructors="", co_organizer="", supervisor="", councilors="",
        subsidy=subsidy, items=_items(raw_items, subsidy), courses=courses, roster=[],
    )
    base.update(overrides)
    return pg.PlanContext(**base)


def _levels(findings, title_part):
    return [f for f in findings if title_part in f.title]


class AllocationTests(unittest.TestCase):
    def test_proportional_integer_and_last_item_absorbs_remainder(self):
        items = _items([
            {"name": "甲", "kind": "一般", "qty": 1, "price": 1, "use": "甲用途"},
            {"name": "乙", "kind": "一般", "qty": 1, "price": 1, "use": "乙用途"},
            {"name": "丙", "kind": "一般", "qty": 1, "price": 1, "use": "丙用途"},
        ], subsidy=10)
        self.assertEqual([i["subsidy"] for i in items], [3, 3, 4])
        self.assertEqual(sum(i["subsidy"] for i in items), 10)
        self.assertEqual([i["self"] for i in items], [-2, -2, -3])  # 補助大於單項，檢查會擋

    def test_remainder_matches_subsidy_when_within_total(self):
        items = _items([
            {"name": "便當", "kind": "餐費", "qty": 60, "price": 100, "use": "用餐"},
            {"name": "棚架", "kind": "一般", "qty": 2, "price": 3500, "use": "休息"},
            {"name": "礦泉水", "kind": "一般", "qty": 10, "price": 150, "use": "飲用"},
            {"name": "救助吊帶", "kind": "一般", "qty": 5, "price": 2200, "use": "確保"},
        ], subsidy=20000)
        self.assertEqual([i["subsidy"] for i in items], [4705, 5490, 1176, 8629])
        self.assertEqual(sum(i["subsidy"] for i in items), 20000)
        self.assertEqual(sum(i["self"] for i in items), 5500)
        self.assertEqual(items[0]["remark"], "用餐")           # 備註只留用途說明，不顯示「○元由…支應」
        for item in items:
            self.assertNotIn("支應", item["remark"])

    def test_nhi_gets_zero_subsidy_and_others_share_the_total(self):
        items = _items([
            {"name": "講義", "kind": "一般", "qty": 10, "price": 100, "use": "上課"},
            {"name": "補充保費", "kind": "二代健保補充保費", "qty": 1, "price": 500, "use": "補充保費"},
            {"name": "場地", "kind": "一般", "qty": 1, "price": 1000, "use": "場地"},
        ], subsidy=1000)
        by_name = {i["name"]: i for i in items}
        self.assertEqual(by_name["補充保費"]["subsidy"], 0)
        self.assertEqual(by_name["補充保費"]["self"], 500)
        self.assertEqual(by_name["補充保費"]["remark"], "補充保費")
        self.assertEqual(by_name["講義"]["subsidy"] + by_name["場地"]["subsidy"], 1000)
        ctx = _ctx(items=[
            {"name": "講義", "kind": "一般", "qty": 10, "price": 100, "use": "上課"},
            {"name": "補充保費", "kind": "二代健保補充保費", "qty": 1, "price": 500, "use": "補充保費"},
            {"name": "場地", "kind": "一般", "qty": 1, "price": 1000, "use": "場地"},
        ], subsidy=1000)
        titles = [f.title for f in pg.check_plan(ctx, today=TODAY)]
        self.assertIn("二代健保補充保費全數由自籌款支應", titles)
        self.assertNotIn("二代健保補充保費不得使用補助款支應", titles)

    def test_meal_cap_uses_allocated_subsidy(self):
        ctx = _ctx(items=[
            {"name": "便當", "kind": "餐費", "qty": 1, "price": 8000, "use": "用餐"},
            {"name": "紀念品", "kind": "紀念品", "qty": 1, "price": 8000, "use": "紀念"},
            {"name": "場地", "kind": "一般", "qty": 1, "price": 4000, "use": "場地"},
        ], subsidy=10000)
        meal = sum(i["subsidy"] for i in ctx.items if i["kind"] in ("餐費", "紀念品"))
        self.assertGreater(meal, 5000)
        hit = _levels(pg.check_plan(ctx, today=TODAY), "餐費及紀念品超過補助金額 50%")
        self.assertEqual(len(hit), 1)
        self.assertEqual(hit[0].level, pg.ERROR)


class CourseTests(unittest.TestCase):
    def test_empty_content_kept_and_count_matches(self):
        start = pg.earliest_activity_date(TODAY)
        courses = [
            {"date": start, "time": "08:10－09:00", "content": "學員簽到", "teacher": ""},
            {"date": start, "time": "09:10－10:00", "content": "", "teacher": ""},
            {"date": start, "time": "10:10－11:00", "content": "", "teacher": ""},
            {"date": start, "time": "12:00－13:30", "content": "午餐及休息", "teacher": ""},
        ]
        ctx = _ctx(courses=courses, start=start, end=start)
        self.assertEqual(len(ctx.courses), 4)
        listed = _levels(pg.check_plan(ctx, today=TODAY), "已列出課程表")
        self.assertIn("共 4 列", listed[0].detail)
        warned = _levels(pg.check_plan(ctx, today=TODAY), "未填課程內容")
        self.assertEqual(warned[0].level, pg.WARN)
        self.assertIn("第 2、3 列", warned[0].detail)

    def test_sorted_by_date_then_time_in_word_and_excel(self):
        day1 = pg.earliest_activity_date(TODAY)
        day2 = day1 + timedelta(days=1)
        courses = [
            {"date": day2, "time": "09:00－10:00", "content": "第二天", "teacher": "乙"},
            {"date": day1, "time": "15:00－16:00", "content": "下午", "teacher": "甲"},
            {"date": day1, "time": "08:10－09:00", "content": "", "teacher": ""},
        ]
        ctx = _ctx(courses=courses, start=day1, end=day2)
        self.assertEqual([c["content"] for c in ctx.courses], ["", "下午", "第二天"])
        self.assertEqual([c["time"] for c in ctx.courses], ["08:10－09:00", "15:00－16:00", "09:00－10:00"])

        from docx import Document
        doc = Document(io.BytesIO(pg.build_docx(ctx)))
        table = doc.tables[0]
        headers = [cell.text for cell in table.rows[0].cells]
        self.assertEqual(headers, ["日期", "時間", "課程內容", "教官／講師"])
        body = [[cell.text for cell in row.cells] for row in table.rows[1:]]
        self.assertEqual(body[0][0], pg.roc_date(day1))
        self.assertEqual(body[0][1], "08:10－09:00")
        self.assertEqual(body[0][2], "")
        self.assertEqual(body[2][0], pg.roc_date(day2))
        self.assertEqual(body[2][2], "第二天")

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(pg.build_xlsx(ctx)))
        ws = wb["課程表"]
        self.assertEqual([ws.cell(1, c).value for c in range(1, 5)], ["日期", "時間", "課程內容", "教官／講師"])
        self.assertEqual(ws["A2"].value, pg.roc_date(day1))
        self.assertEqual(ws["B2"].value, "08:10－09:00")
        self.assertEqual(ws["C4"].value, "第二天")
        self.assertEqual(wb["經費概算表"]["G3"].value, "供訓練人員中餐及晚餐")

    def test_missing_date_defaults_to_start(self):
        start = pg.earliest_activity_date(TODAY)
        ctx = _ctx(courses=[{"time": "09:00－10:00", "content": "課程", "teacher": "教官"}], start=start, end=start)
        self.assertEqual(ctx.courses[0]["date"], start)

    def test_date_outside_range_is_error(self):
        start = pg.earliest_activity_date(TODAY)
        ctx = _ctx(
            courses=[{"date": start + timedelta(days=3), "time": "09:00－10:00", "content": "課程", "teacher": ""}],
            start=start, end=start,
        )
        hit = _levels(pg.check_plan(ctx, today=TODAY), "課程日期不在活動期間")
        self.assertEqual(hit[0].level, pg.ERROR)


class WordingTests(unittest.TestCase):
    def test_councilor_suffix_not_doubled(self):
        self.assertEqual(pg.format_councilors("王議員"), "王議員")
        self.assertEqual(pg.format_councilors("王議員、李議員"), "王議員、李議員")
        self.assertEqual(pg.format_councilors("王、李"), "王議員、李議員")
        self.assertEqual(pg.format_councilors(""), "ＯＯ議員")
        self.assertEqual(pg.format_councilors("議員"), "ＯＯ議員")
        ctx = _ctx(councilors="王議員")
        text = pg.fund_text(ctx)
        self.assertIn("王議員建議", text)
        self.assertNotIn("王議員議員", text)

    def test_year_in_activity_name_is_stripped(self):
        ctx = _ctx(name="一一五年度繩索救助訓練", year=115)
        self.assertEqual(ctx.title2, "115年度繩索救助訓練活動計畫")
        self.assertNotIn("一一五年度", ctx.title2)
        ctx2 = _ctx(name="115年度一一五年度繩索救助訓練", year=115)
        self.assertEqual(ctx2.title2, "115年度繩索救助訓練活動計畫")
        warned = _levels(pg.check_plan(ctx, today=TODAY), "活動名稱含有年度")
        self.assertEqual(warned[0].level, pg.WARN)
        only_year = _ctx(name="一一五年度", year=115)
        self.assertTrue(_levels(pg.check_plan(only_year, today=TODAY), "活動名稱只寫了年度"))

    def test_organizer_and_co_organizer_both_appear(self):
        ctx = _ctx(co_organizer="大武大隊達仁分隊")
        org = dict(pg.sections(ctx))["辦理單位"]
        self.assertTrue(any(line.startswith("三、承辦單位：達仁義消分隊") for line in org))
        self.assertTrue(any(line.startswith("四、協辦單位：大武大隊達仁分隊") for line in org))
        ctx2 = _ctx(co_organizer="")
        org2 = dict(pg.sections(ctx2))["辦理單位"]
        self.assertTrue(any("承辦單位：達仁義消分隊" in line for line in org2))
        self.assertFalse(any("協辦單位" in line for line in org2))

    def test_gift_item_warns(self):
        ctx = _ctx(items=[
            {"name": "贈品", "kind": "一般", "qty": 10, "price": 50, "use": "發給學員"},
            {"name": "場地", "kind": "一般", "qty": 1, "price": 30000, "use": "場地"},
        ], subsidy=10000)
        hit = _levels(pg.check_plan(ctx, today=TODAY), "贈品")
        self.assertEqual(hit[0].level, pg.WARN)
        self.assertIn("宣導品", hit[0].detail)
        self.assertIn("獎品", hit[0].detail)

    def test_zero_and_negative_qty_or_price_error_and_kept(self):
        rows = [
            {"name": "空白數量", "kind": "一般", "qty": 0, "price": 100, "use": "測試"},
            {"name": "負單價", "kind": "一般", "qty": 2, "price": -50, "use": "測試"},
            {"name": "正常", "kind": "一般", "qty": 1, "price": 1000, "use": "測試"},
        ]
        items = _items(rows, subsidy=500)
        self.assertEqual([i["name"] for i in items], ["空白數量", "負單價", "正常"])
        ctx = _ctx(items=rows, subsidy=500)
        hit = _levels(pg.check_plan(ctx, today=TODAY), "數量或單價不是正數")
        self.assertEqual(hit[0].level, pg.ERROR)
        self.assertIn("空白數量", hit[0].detail)
        self.assertIn("負單價", hit[0].detail)
        self.assertEqual(items[2]["subsidy"], 500)
        self.assertEqual(items[0]["subsidy"], 0)
        self.assertEqual(items[1]["subsidy"], 0)


class DateRuleTests(unittest.TestCase):
    def test_threshold_is_one_calendar_month_plus_seven_days(self):
        self.assertEqual(pg.earliest_activity_date(date(2026, 10, 22)), date(2026, 11, 29))
        self.assertEqual(pg.earliest_activity_date(date(2026, 1, 31)), date(2026, 3, 7))
        self.assertEqual(pg.add_months(date(2024, 1, 31), 1), date(2024, 2, 29))

    def test_past_date_is_error_without_negative_day_count(self):
        start = TODAY - timedelta(days=21)
        ctx = _ctx(start=start, end=start)
        hit = _levels(pg.check_plan(ctx, today=TODAY), "活動開始日期已過")
        self.assertEqual(hit[0].level, pg.ERROR)
        self.assertNotIn("只剩 -", hit[0].title)
        self.assertNotIn("只剩 -", hit[0].detail)
        self.assertIn("21", hit[0].detail)

    def test_short_lead_warns_with_month_plus_week(self):
        earliest = pg.earliest_activity_date(TODAY)
        ctx = _ctx(start=earliest - timedelta(days=1), end=earliest - timedelta(days=1))
        hit = _levels(pg.check_plan(ctx, today=TODAY), "不足 1 個月又 1 週")
        self.assertEqual(hit[0].level, pg.WARN)
        self.assertIn("1 個月再加 7 天", hit[0].detail)
        self.assertNotIn("不足 1 個月」", hit[0].title)

    def test_deadline_is_activity_minus_one_month_minus_seven_days(self):
        self.assertEqual(pg.submission_deadline(date(2026, 11, 20)), date(2026, 10, 13))
        self.assertEqual(pg.submission_deadline(date(2026, 12, 31)), date(2026, 11, 23))
        late = _ctx(start=date(2026, 11, 20), end=date(2026, 11, 20))
        hit = _levels(pg.check_plan(late, today=TODAY), "已超過送件期限")
        self.assertEqual(hit[0].level, pg.WARN)
        self.assertIn("115年10月13日", hit[0].detail)
        ok = _ctx(start=date(2026, 12, 31), end=date(2026, 12, 31))
        passed = _levels(pg.check_plan(ok, today=TODAY), "已預留 1 個月又 1 週")
        self.assertEqual(passed[0].level, pg.OK)
        self.assertIn("115年11月23日", passed[0].detail)

    def test_lecturer_meal_and_nhi_allocation(self):
        items = _items([
            {"name": "講師鐘點費", "kind": "內聘鐘點費", "qty": 1, "price": 8000, "use": "授課"},
            {"name": "膳費", "kind": "餐費", "qty": 1, "price": 6000, "use": "用餐"},
            {"name": "二代健保補充保費", "kind": "二代健保補充保費", "qty": 1, "price": 1000, "use": "補充保費"},
        ], subsidy=10000)
        by_name = {item["name"]: item for item in items}
        self.assertEqual(by_name["講師鐘點費"]["subsidy"], 5714)
        self.assertEqual(by_name["膳費"]["subsidy"], 4286)
        self.assertEqual(by_name["二代健保補充保費"]["subsidy"], 0)
        self.assertEqual(by_name["二代健保補充保費"]["self"], 1000)
        self.assertEqual(sum(item["subsidy"] for item in items), 10000)

    def test_exact_earliest_date_passes(self):
        earliest = pg.earliest_activity_date(TODAY)
        ctx = _ctx(start=earliest, end=earliest)
        hit = _levels(pg.check_plan(ctx, today=TODAY), "已預留 1 個月又 1 週")
        self.assertEqual(hit[0].level, pg.OK)

    def test_end_before_start_is_error(self):
        start = pg.earliest_activity_date(TODAY)
        ctx = _ctx(start=start, end=start - timedelta(days=1))
        hit = _levels(pg.check_plan(ctx, today=TODAY), "結束日期早於開始日期")
        self.assertEqual(hit[0].level, pg.ERROR)

    def test_subsidy_over_total_does_not_show_negative_self_fund(self):
        ctx = _ctx(items=[
            {"name": "場地", "kind": "一般", "qty": 1, "price": 10000, "use": "場地"},
        ], subsidy=12000)
        self.assertLess(ctx.self_fund, 0)
        hit = _levels(pg.check_plan(ctx, today=TODAY), "補助金額大於總經費")
        self.assertEqual(hit[0].level, pg.ERROR)
        note = pg.budget_note(ctx)
        self.assertNotIn("-", note)
        self.assertIn("補助金額超過總經費", note)
        texts = " ".join(f"{f.title} {f.detail}" for f in pg.check_plan(ctx, today=TODAY))
        self.assertNotIn("-", texts.replace("1 個月又 1 週", "").replace("—", ""))


class ExampleTests(unittest.TestCase):
    def test_example_fills_required_fields_and_passes(self):
        data = pg.example_inputs(TODAY)
        for key in ("unit", "name", "place", "audience"):
            self.assertTrue(str(data[key]).strip(), key)
        self.assertGreaterEqual(data["start"], pg.earliest_activity_date(TODAY))
        ctx = pg.example_context(TODAY)
        findings = pg.check_plan(ctx, today=TODAY)
        errors = [f for f in findings if f.level == pg.ERROR]
        self.assertEqual(errors, [], [f"{f.title}：{f.detail}" for f in errors])
        self.assertIn("王議員建議", pg.fund_text(ctx))
        self.assertNotIn("王議員議員", pg.fund_text(ctx))
        docx_bytes = pg.build_docx(ctx)
        xlsx_bytes = pg.build_xlsx(ctx)
        self.assertTrue(docx_bytes.startswith(b"PK"))
        self.assertTrue(xlsx_bytes.startswith(b"PK"))


class ComplianceReportTests(unittest.TestCase):
    def test_failures_name_rule_location_comparison_and_suggestion(self):
        ctx = _ctx(
            start=date(2026, 11, 20),
            end=date(2026, 11, 20),
            items=[
                {"name": "膳費", "kind": "餐費", "qty": 1, "unit": "份", "price": 8000, "use": "用餐"},
                {"name": "雜支", "kind": "雜支", "qty": 1, "unit": "式", "price": 2000, "use": "雜項"},
            ],
            subsidy=8000,
            courses=[{
                "date": date(2026, 11, 20), "time": "09:00－12:00",
                "content": "訓練課程", "teacher": "教官",
            }],
        )
        findings = pg.check_plan(ctx, today=TODAY)
        problems = [f for f in findings if f.level != pg.OK]
        self.assertGreaterEqual(len(problems), 2)
        for finding in problems:
            self.assertTrue(finding.rule, finding.title)
            self.assertTrue(finding.where, finding.title)
            self.assertTrue(finding.comparison, finding.title)
            self.assertTrue(finding.suggestion, finding.title)
            text = pg.format_compliance_item(finding)
            self.assertIn("規則：", text)
            self.assertIn("位置：", text)
            self.assertIn("現況與規定：", text)
            self.assertIn("修正建議：", text)

        late = _levels(findings, "已超過送件期限")[0]
        self.assertEqual(late.level, pg.WARN)
        self.assertIn("活動日期", late.where)
        self.assertIn("錯誤樣態「活動類案件未於活動 1 個月前送達局本部」", late.rule)
        self.assertIn("注意要點「申請流程與發文單位」", late.rule)
        self.assertIn("屆時送局期限 2026-10-13，今日 2026-10-22 已逾 9 天", late.comparison)

        meal = _levels(findings, "餐費及紀念品超過補助金額 50%")[0]
        self.assertEqual(meal.level, pg.ERROR)
        self.assertIn("經費概算表「膳費」", meal.where)
        self.assertIn("注意要點「經費比例上限（自籌款、餐費紀念品、雜支）」", meal.rule)
        self.assertIn("錯誤樣態「餐費及紀念品超過申請補助金額 50%", meal.rule)
        self.assertIn("上限", meal.comparison)

        misc = _levels(findings, "雜支超過總經費 5%")[0]
        self.assertEqual(misc.level, pg.ERROR)
        self.assertIn("經費概算表「雜支」", misc.where)

        kind, banner = pg.compliance_banner(findings)
        self.assertEqual(kind, "error")
        self.assertNotEqual(banner, pg.ALL_CLEAR_MESSAGE)

    def test_all_checks_pass_uses_exact_success_wording(self):
        findings = pg.check_plan(pg.example_context(TODAY), today=TODAY)
        self.assertTrue(findings)
        self.assertEqual(
            [f.title for f in findings if f.level != pg.OK],
            [],
        )
        kind, banner = pg.compliance_banner(findings)
        self.assertEqual(kind, "pass")
        self.assertEqual(banner, "計畫符合現行補捐助規定，可以開始匯出")
        self.assertEqual(banner, pg.ALL_CLEAR_MESSAGE)


class AppLoaderTests(unittest.TestCase):
    def test_app_loader_calls_normalize_items_with_subsidy(self):
        import inspect
        import app

        pg_from_app = app.load_plan_generator()
        self.assertIn("subsidy", inspect.signature(pg_from_app.normalize_items).parameters)
        items = pg_from_app.normalize_items(
            [
                {"name": "講師鐘點費", "kind": "內聘鐘點費", "qty": 1, "price": 8000, "use": "授課"},
                {"name": "膳費", "kind": "餐費", "qty": 1, "price": 6000, "use": "用餐"},
                {"name": "二代健保補充保費", "kind": "二代健保補充保費", "qty": 1, "price": 1000, "use": "補充保費"},
            ],
            subsidy=10000,
            source_label="議員建議補助款",
        )
        by_name = {item["name"]: item for item in items}
        self.assertEqual(by_name["講師鐘點費"]["subsidy"], 5714)
        self.assertEqual(by_name["膳費"]["subsidy"], 4286)
        self.assertEqual(by_name["二代健保補充保費"]["subsidy"], 0)

    def test_stale_module_without_subsidy_is_reloaded(self):
        import app
        import plan_generator as live

        def stale_normalize(rows, source_label="議員建議補助款"):
            return []

        live.normalize_items = stale_normalize
        live.MODULE_REVISION = 0
        try:
            loaded = app.load_plan_generator()
            items = loaded.normalize_items(
                [{"name": "場地", "kind": "一般", "qty": 1, "price": 1000, "use": "場地"}],
                subsidy=1000,
                source_label="議員建議補助款",
            )
            self.assertEqual(items[0]["subsidy"], 1000)
            self.assertEqual(loaded.MODULE_REVISION, 9)
        finally:
            import importlib
            importlib.reload(live)


if __name__ == "__main__":
    unittest.main()
